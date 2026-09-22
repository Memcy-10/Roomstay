from datetime import date, datetime
from io import BytesIO
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc

from app.core.db import get_db
from app.core.config import settings
from app.core.responses import success_response
from app.dependencies.auth import get_current_user, require_host_or_admin
from app.models.user import User
from app.models.room import Room
from app.models.sale import Sale, SaleDetail, SaleStatus
from app.models.invoice import Invoice, InvoiceDetail, InvoiceStatus

router = APIRouter(prefix="/reports")


@router.get("/daily-sales")
def daily_sales_report(
    fecha: Optional[date] = Query(None),
    current_user: User = Depends(require_host_or_admin),
    db: Session = Depends(get_db),
):
    if not fecha:
        fecha = date.today()

    query = db.query(Sale).filter(Sale.fechaVenta == fecha)
    if current_user.role == "host":
        query = query.filter(Sale.hostId == current_user.id)

    ventas = query.order_by(Sale.createdAt.asc()).all()

    user_ids = {s.userId for s in ventas if s.userId}
    users_map = {u.id: u for u in db.query(User).filter(User.id.in_(list(user_ids))).all()}

    data = []
    total_general = 0.0
    for s in ventas:
        u = users_map.get(s.userId)
        detalles = []
        if s.details:
            for d in s.details:
                detalles.append({
                    "descripcion": d.descripcion,
                    "cantidad": d.cantidad,
                    "precioUnitario": float(d.precioUnitario),
                    "subtotal": float(d.subtotal),
                    "total": float(d.total),
                })
        total = float(s.total)
        total_general += total
        data.append({
            "id": s.id,
            "numeroVenta": s.numeroVenta,
            "fecha": str(s.fechaVenta),
            "cliente": f"{u.firstName} {u.lastName}" if u else "N/A",
            "clienteDocumento": u.documentNumber if u else "N/A",
            "clienteEmail": u.email if u else "N/A",
            "estado": s.estado.value if isinstance(s.estado, SaleStatus) else s.estado,
            "metodoPago": s.metodoPago,
            "subTotal": float(s.subTotal),
            "descuento": float(s.descuento),
            "impuestos": float(s.impuestos),
            "total": total,
            "detalles": detalles,
        })

    return success_response(
        data={
            "fecha": str(fecha),
            "ventas": data,
            "totalVentas": len(data),
            "totalGeneral": total_general,
            "empresa": {
                "nombre": settings.COMPANY_NAME,
                "nit": settings.COMPANY_NIT,
                "direccion": settings.COMPANY_ADDRESS,
                "email": settings.COMPANY_EMAIL,
                "telefono": settings.COMPANY_PHONE,
            },
            "generado": datetime.now().isoformat(timespec="seconds"),
        },
        message="Reporte diario obtenido"
    )


def _generar_pdf_reporte_ventas(report_data: dict) -> BytesIO:
    buffer = BytesIO()
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.5*inch, rightMargin=0.5*inch,
                                topMargin=0.5*inch, bottomMargin=0.5*inch)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('Title2', parent=styles['Title'], fontSize=18, alignment=1, textColor=colors.HexColor('#0f766e'))
        hdr_style = ParagraphStyle('Hdr', parent=styles['Normal'], fontSize=9, textColor=colors.grey)
        emp = report_data.get("empresa", {})

        story = []
        story.append(Paragraph(emp.get("nombre", "EMPRESA"), title_style))
        story.append(Paragraph(f"NIT: {emp.get('nit','')} | Dirección: {emp.get('direccion','')}", hdr_style))
        story.append(Paragraph(f"Tel: {emp.get('telefono','')} | Email: {emp.get('email','')}", hdr_style))
        story.append(Spacer(1, 0.15*inch))
        story.append(Paragraph("REPORTE DIARIO DE VENTAS", styles['Heading2']))
        story.append(Paragraph(f"Fecha: {report_data.get('fecha')} &nbsp;&nbsp;|&nbsp;&nbsp; Generado: {report_data.get('generado','')}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))

        ventas = report_data.get("ventas", [])
        if not ventas:
            story.append(Paragraph("No se encontraron ventas para la fecha seleccionada.", styles['Normal']))
        else:
            tbl_data = [["N° Venta", "Cliente", "Documento", "Estado", "Método Pago", "Subtotal", "Desc.", "Impuestos", "Total"]]
            for v in ventas:
                tbl_data.append([
                    str(v.get("numeroVenta","")),
                    str(v.get("cliente",""))[:30],
                    str(v.get("clienteDocumento","")),
                    str(v.get("estado","")),
                    str(v.get("metodoPago","")),
                    f"{v.get('subTotal',0):,.0f}",
                    f"{v.get('descuento',0):,.0f}",
                    f"{v.get('impuestos',0):,.0f}",
                    f"{v.get('total',0):,.0f}",
                ])
            tbl_data.append(["", "", "", "", "", "TOTAL GENERAL", "", "", f"{report_data.get('totalGeneral',0):,.0f}"])

            tbl = Table(tbl_data, colWidths=[1.0*inch, 1.6*inch, 0.9*inch, 0.8*inch, 1.0*inch, 0.8*inch, 0.6*inch, 0.7*inch, 0.9*inch])
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('ALIGN', (-1,0), (-1,-1), 'RIGHT'),
                ('ALIGN', (0,0), (0,-1), 'LEFT'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#f0fdfa')]),
                ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#ccfbf1')),
                ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
            ]))
            story.append(tbl)

            story.append(Spacer(1, 0.3*inch))
            story.append(Paragraph("DETALLE POR VENTA", styles['Heading3']))
            story.append(Spacer(1, 0.1*inch))
            for v in ventas:
                story.append(Paragraph(f"<b>{v.get('numeroVenta','')} - {v.get('cliente','')}</b>", styles['Normal']))
                if v.get("detalles"):
                    d_data = [["Descripción", "Cantidad", "Precio Unitario", "Subtotal", "Total"]]
                    for d in v["detalles"]:
                        d_data.append([
                            str(d.get("descripcion",""))[:60],
                            str(d.get("cantidad",1)),
                            f"{d.get('precioUnitario',0):,.0f}",
                            f"{d.get('subtotal',0):,.0f}",
                            f"{d.get('total',0):,.0f}",
                        ])
                    dt = Table(d_data, colWidths=[3.2*inch, 0.9*inch, 1.3*inch, 1.2*inch, 1.2*inch])
                    dt.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e2e8f0')),
                        ('FONTSIZE', (0,0), (-1,-1), 8),
                        ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
                    ]))
                    story.append(dt)
                    story.append(Spacer(1, 0.15*inch))

        doc.build(story)
        buffer.seek(0)
    except Exception as e:
        buffer.close()
        raise e
    return buffer


def _generar_excel_reporte_ventas(report_data: dict) -> BytesIO:
    buffer = BytesIO()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"

        emp = report_data.get("empresa", {})
        ws["A1"] = emp.get("nombre", "EMPRESA")
        ws["A1"].font = Font(bold=True, size=16, color="0F766E")
        ws.merge_cells("A1:I1")

        ws["A2"] = f"NIT: {emp.get('nit','')} | Dirección: {emp.get('direccion','')}"
        ws.merge_cells("A2:I2")
        ws["A3"] = f"Tel: {emp.get('telefono','')} | Email: {emp.get('email','')}"
        ws.merge_cells("A3:I3")

        ws["A5"] = "REPORTE DIARIO DE VENTAS"
        ws["A5"].font = Font(bold=True, size=14)
        ws.merge_cells("A5:I5")
        ws["A6"] = f"Fecha: {report_data.get('fecha')} | Generado: {report_data.get('generado','')}"
        ws.merge_cells("A6:I6")

        hdr_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        hdr_font = Font(bold=True, color="FFFFFF")

        headers = ["N° Venta", "Cliente", "Documento", "Email", "Estado", "Método Pago", "Subtotal", "Descuento", "Impuestos", "Total"]
        row = 8
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=col, value=h)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal="center")

        ventas = report_data.get("ventas", [])
        for i, v in enumerate(ventas, 1):
            r = row + i
            ws.cell(row=r, column=1, value=v.get("numeroVenta",""))
            ws.cell(row=r, column=2, value=v.get("cliente",""))
            ws.cell(row=r, column=3, value=v.get("clienteDocumento",""))
            ws.cell(row=r, column=4, value=v.get("clienteEmail",""))
            ws.cell(row=r, column=5, value=v.get("estado",""))
            ws.cell(row=r, column=6, value=v.get("metodoPago",""))
            ws.cell(row=r, column=7, value=v.get("subTotal",0))
            ws.cell(row=r, column=8, value=v.get("descuento",0))
            ws.cell(row=r, column=9, value=v.get("impuestos",0))
            ws.cell(row=r, column=10, value=v.get("total",0))

        total_row = row + len(ventas) + 1
        ws.cell(row=total_row, column=5, value="TOTAL GENERAL").font = Font(bold=True)
        ws.cell(row=total_row, column=10, value=report_data.get("totalGeneral",0)).font = Font(bold=True)

        # Detalle sheet
        ws2 = wb.create_sheet("Detalle")
        ws2["A1"] = "DETALLE DE VENTAS"
        ws2["A1"].font = Font(bold=True, size=14)
        det_hdr = ["N° Venta", "Descripción", "Cantidad", "Precio Unitario", "Subtotal", "Total"]
        for col, h in enumerate(det_hdr, 1):
            c = ws2.cell(row=3, column=col, value=h)
            c.fill = hdr_fill
            c.font = hdr_font
        dr = 4
        for v in ventas:
            if v.get("detalles"):
                for d in v["detalles"]:
                    ws2.cell(row=dr, column=1, value=v.get("numeroVenta",""))
                    ws2.cell(row=dr, column=2, value=d.get("descripcion",""))
                    ws2.cell(row=dr, column=3, value=d.get("cantidad",1))
                    ws2.cell(row=dr, column=4, value=d.get("precioUnitario",0))
                    ws2.cell(row=dr, column=5, value=d.get("subtotal",0))
                    ws2.cell(row=dr, column=6, value=d.get("total",0))
                    dr += 1

        for col_letter in ["A","B","C","D","E","F","G","H","I","J"]:
            try:
                ws.column_dimensions[col_letter].width = 18
            except Exception:
                pass

        wb.save(buffer)
        buffer.seek(0)
    except Exception as e:
        buffer.close()
        raise e
    return buffer


@router.get("/daily-sales/pdf")
def daily_sales_pdf(
    fecha: Optional[date] = Query(None),
    current_user: User = Depends(require_host_or_admin),
    db: Session = Depends(get_db),
):
    import json
    report_resp = daily_sales_report(fecha=fecha, current_user=current_user, db=db)
    report_data = json.loads(report_resp.body)
    data = report_data.get("data", report_data)
    try:
        buffer = _generar_pdf_reporte_ventas(data)
    except ImportError as ie:
        raise HTTPException(status_code=500, detail=f"Librería no disponible: {ie}. Instale reportlab.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando PDF: {e}")

    fname = f"reporte_ventas_{fecha or date.today()}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'}
    )


@router.get("/daily-sales/excel")
def daily_sales_excel(
    fecha: Optional[date] = Query(None),
    current_user: User = Depends(require_host_or_admin),
    db: Session = Depends(get_db),
):
    import json
    report_resp = daily_sales_report(fecha=fecha, current_user=current_user, db=db)
    report_data = json.loads(report_resp.body)
    data = report_data.get("data", report_data)
    try:
        buffer = _generar_excel_reporte_ventas(data)
    except ImportError as ie:
        raise HTTPException(status_code=500, detail=f"Librería no disponible: {ie}. Instale openpyxl.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando Excel: {e}")

    fname = f"reporte_ventas_{fecha or date.today()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'}
    )


def _generar_pdf_factura(invoice_data: dict, sale_details: list, company: dict) -> BytesIO:
    buffer = BytesIO()
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.6*inch, rightMargin=0.6*inch,
                                topMargin=0.6*inch, bottomMargin=0.6*inch)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph(company.get("nombre", "EMPRESA"), ParagraphStyle('Emp', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#0f766e'))))
        story.append(Paragraph(f"NIT: {company.get('nit','')}", styles['Normal']))
        story.append(Paragraph(f"Dirección: {company.get('direccion','')}", styles['Normal']))
        story.append(Paragraph(f"Tel: {company.get('telefono','')} | Email: {company.get('email','')}", styles['Normal']))
        story.append(Spacer(1, 0.25*inch))

        info_tbl = Table([
            [Paragraph(f"<b>FACTURA</b>", styles['Heading2']), f"Fecha emisión: {invoice_data.get('fechaEmision','')}"],
            [f"N° Factura: <b>{invoice_data.get('numeroFactura','')}</b>", f"Vence: {invoice_data.get('fechaVencimiento','') or 'N/A'}"],
            [f"Venta asociada: {invoice_data.get('numeroVenta','')}", f"Estado: {invoice_data.get('estado','')}"],
        ], colWidths=[3.5*inch, 3.5*inch])
        info_tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#0f766e')),
            ('BACKGROUND', (0,0), (0,0), colors.HexColor('#f0fdfa')),
            ('SPAN', (0,0), (0,0)),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 10),
            ('TOPPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(info_tbl)
        story.append(Spacer(1, 0.25*inch))

        story.append(Paragraph("<b>DATOS DEL CLIENTE</b>", styles['Heading3']))
        cliente_tbl = Table([
            [f"Nombre: {invoice_data.get('userFirstName','')} {invoice_data.get('userLastName','')}",
             f"Tipo Doc.: {invoice_data.get('userDocumentType','')}"],
            [f"Documento: {invoice_data.get('userDocument','')}",
             f"Teléfono: {invoice_data.get('userPhone','')}"],
            [f"Email: {invoice_data.get('userEmail','')}",
             f"Dirección: {invoice_data.get('userAddress','')}"],
        ], colWidths=[3.5*inch, 3.5*inch])
        cliente_tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.grey),
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#fafafa')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(cliente_tbl)
        story.append(Spacer(1, 0.25*inch))

        story.append(Paragraph("<b>DETALLE</b>", styles['Heading3']))
        det_data = [["Item", "Descripción", "Cantidad", "Precio U.", "Subtotal", "Desc.", "Impuesto", "Total"]]
        for i, d in enumerate(sale_details or [], 1):
            det_data.append([
                str(i), str(d.get("descripcion",""))[:50],
                str(d.get("cantidad",1)),
                f"{float(d.get('precioUnitario',0)):,.0f}",
                f"{float(d.get('subtotal',0)):,.0f}",
                f"{float(d.get('descuento',0)):,.0f}",
                f"{float(d.get('impuesto',0)):,.0f}",
                f"{float(d.get('total',0)):,.0f}",
            ])
        det_data.append(["", "", "", "", "SUBTOTAL", "", "", f"{float(invoice_data.get('subTotal',0)):,.0f}"])
        det_data.append(["", "", "", "", "DESCUENTO", "", "", f"{float(invoice_data.get('descuento',0)):,.0f}"])
        det_data.append(["", "", "", "", "IMPUESTOS", "", "", f"{float(invoice_data.get('impuestos',0)):,.0f}"])
        det_data.append(["", "", "", "", "TOTAL", "", "", f"{float(invoice_data.get('total',0)):,.0f}"])

        dt = Table(det_data, colWidths=[0.5*inch, 2.8*inch, 0.8*inch, 1.0*inch, 0.9*inch, 0.7*inch, 0.7*inch, 0.9*inch])
        dt.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-5), 9),
            ('GRID', (0,0), (-1,-5), 0.5, colors.grey),
            ('ALIGN', (2,0), (-1,-1), 'RIGHT'),
            ('BACKGROUND', (-5,-4), (-1,-1), colors.HexColor('#ccfbf1')),
            ('FONTNAME', (-5,-1), (-1,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (-5,-4), (-1,-1), 10),
        ]))
        story.append(dt)

        story.append(Spacer(1, 0.5*inch))
        if invoice_data.get("notas"):
            story.append(Paragraph(f"<b>Notas:</b> {invoice_data['notas']}", styles['Normal']))
            story.append(Spacer(1, 0.3*inch))

        firma_tbl = Table([
            ["____________________________________", "____________________________________"],
            ["Firma autorizada", "Cliente"],
        ], colWidths=[3.5*inch, 3.5*inch])
        firma_tbl.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(firma_tbl)

        doc.build(story)
        buffer.seek(0)
    except Exception as e:
        buffer.close()
        raise e
    return buffer


@router.get("/invoice/{invoice_id}/pdf")
def invoice_pdf(
    invoice_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    invoice = db.query(Invoice).filter(Invoice.id == invoice_id).first()
    if not invoice:
        raise HTTPException(status_code=404, detail="Factura no encontrada")

    sale = db.query(Sale).filter(Sale.id == invoice.saleId).first()
    is_admin = current_user.role == "admin"
    is_owner = invoice.userId == current_user.id
    is_host = sale is not None and sale.hostId == current_user.id
    if not (is_admin or is_owner or is_host):
        raise HTTPException(status_code=403, detail="No autorizado")

    user = db.query(User).filter(User.id == invoice.userId).first()

    sale_details = []
    if invoice.details:
        for d in invoice.details:
            sale_details.append({
                "descripcion": d.descripcion, "cantidad": d.cantidad,
                "precioUnitario": float(d.precioUnitario), "subtotal": float(d.subtotal),
                "descuento": float(d.descuento), "impuesto": float(d.impuesto),
                "total": float(d.total),
            })
    elif sale and sale.details:
        for d in sale.details:
            sale_details.append({
                "descripcion": d.descripcion, "cantidad": d.cantidad,
                "precioUnitario": float(d.precioUnitario), "subtotal": float(d.subtotal),
                "descuento": float(d.descuento), "impuesto": float(d.impuesto),
                "total": float(d.total),
            })

    invoice_data = {
        "numeroFactura": invoice.numeroFactura,
        "numeroVenta": sale.numeroVenta if sale else "",
        "fechaEmision": str(invoice.fechaEmision),
        "fechaVencimiento": str(invoice.fechaVencimiento) if invoice.fechaVencimiento else "",
        "subTotal": float(invoice.subTotal),
        "descuento": float(invoice.descuento),
        "impuestos": float(invoice.impuestos),
        "total": float(invoice.total),
        "estado": invoice.estado.value if isinstance(invoice.estado, InvoiceStatus) else invoice.estado,
        "notas": invoice.notas,
        "userFirstName": user.firstName if user else "",
        "userLastName": user.lastName if user else "",
        "userDocumentType": getattr(user, 'documentType', '') if user else "",
        "userDocument": getattr(user, 'documentNumber', '') if user else "",
        "userEmail": user.email if user else "",
        "userPhone": user.phone if user else "",
        "userAddress": getattr(user, 'address', '') if user else "",
    }

    company = {
        "nombre": settings.COMPANY_NAME, "nit": settings.COMPANY_NIT,
        "direccion": settings.COMPANY_ADDRESS, "email": settings.COMPANY_EMAIL,
        "telefono": settings.COMPANY_PHONE,
    }

    try:
        buffer = _generar_pdf_factura(invoice_data, sale_details, company)
    except ImportError as ie:
        raise HTTPException(status_code=500, detail=f"Librería no disponible: {ie}. Instale reportlab.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando PDF factura: {e}")

    fname = f"factura_{invoice.numeroFactura}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{fname}"'}
    )
